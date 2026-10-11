# -*- coding: utf-8 -*-
"""Генерация отчёта по итоговому проекту FitShop."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(__file__).resolve().parents[2] / 'П-6-23 УП.04.01 Давтян Г.Л. FitShop.docx'
# Prefer project parent (УП итоговый проект)
OUT = Path(r'C:\Users\gorda\OneDrive\Desktop\УП итоговый проект') / 'П-6-23 УП.04.01 Давтян Г.Л. FitShop.docx'


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            table.rows[r_i + 1].cells[c_i].text = str(val)
    doc.add_paragraph()


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)

    # Title
    for line in [
        'МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ',
        'ФГБОУ ВО «Российский экономический университет имени Г.В. Плеханова»',
        'Московский приборостроительный техникум',
    ]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run('ОТЧЁТ')
    r.bold = True
    r.font.size = Pt(16)

    add_para(doc, 'по учебной практике УП.04.01')
    add_para(doc, 'Индивидуальный итоговый проект')
    add_para(doc, 'Тема: интернет-магазин фитнес-продукции (FitShop)')
    add_para(doc, 'Студент: Давтян Гор Левонович, группа П-6-23')
    add_para(doc, 'Руководитель: Горбутова Маргарита Витальевна')
    add_para(doc, '2026 год')
    doc.add_page_break()

    add_heading(doc, 'Содержание', 1)
    for item in [
        '1. Тема индивидуального проекта',
        '2. Описание предметной области',
        '3. Диаграмма базы данных',
        '4. Словарь данных',
        '5. Экспорт схемы / скрипт инициализации',
        '6. Контракт API',
        '7. Архитектура клиента',
        '8. Ключевые фрагменты реализации',
        '9. Демонстрация работы',
        '10. Перечень тестов',
        '11. Вывод',
    ]:
        add_para(doc, item)

    add_heading(doc, '1. Тема индивидуального проекта', 1)
    add_para(
        doc,
        'Разработка web-клиента интернет-магазина фитнес-продукции (спортивное питание '
        'и аксессуары) на Flutter Web с серверной частью на платформе PocketBase.',
    )

    add_heading(doc, '2. Описание предметной области', 1)
    add_para(
        doc,
        'Система обслуживает три роли пользователей. Покупатель просматривает каталог, '
        'оформляет заказ с автоматическим расчётом стоимости (акции, карта лояльности, '
        'доставка) и видит свои заказы. Менеджер ведёт справочники и каталог, управляет '
        'акциями и меняет статусы заказов на экране «Сборка». Администратор управляет '
        'пользователями, смотрит сводную аналитику и выполняет физическое удаление / '
        'восстановление записей.',
    )
    add_para(doc, 'Правила ценообразования (содержательная особенность проекта):', bold=True)
    add_para(doc, '— скидка по активной акции на товар (берётся максимальный процент);')
    add_para(doc, '— скидка лояльности: bronze 3%, silver 7%, gold 12% от суммы после акций;')
    add_para(doc, '— бесплатная доставка при сумме товара от 5000 ₽, иначе 350 ₽.')

    add_heading(doc, '3. Диаграмма базы данных', 1)
    add_para(doc, 'Логическая схема (сущности и связи):')
    add_para(
        doc,
        'Brand, Supplier, Category — справочники; Product связан с Brand и Supplier (M:1), '
        'с Category (M:N). Customer связан с LoyaltyCard (1:1) и Order (1:N). Order содержит '
        'OrderItem (1:N). Promotion связан с Product (M:N). Review связан с Product и Customer. '
        'Users связан с Customer.',
    )
    add_para(doc, 'Типы связей:')
    add_para(doc, '— 1:1: Customer ↔ LoyaltyCard;')
    add_para(doc, '— 1:N: Brand/Supplier → Product, Customer → Order, Order → OrderItem;')
    add_para(doc, '— M:N: Product ↔ Category, Product ↔ Promotion.')
    add_para(
        doc,
        'Физическая схема реализована коллекциями PocketBase (SQLite). Полное описание полей — '
        'в файле pocketbase/схема.md; машинный экспорт — pocketbase/pb_schema_export.json.',
    )

    add_heading(doc, '4. Словарь данных', 1)
    add_table(
        doc,
        ['Сущность', 'Поле', 'Тип', 'Ограничения', 'Назначение'],
        [
            ['products', 'name', 'text', 'обязательное', 'Название товара'],
            ['products', 'price', 'number', '≥ 0', 'Цена'],
            ['products', 'stock', 'number', 'целое ≥ 0', 'Остаток'],
            ['products', 'brand', 'relation', '→ brands', 'Бренд (M:1)'],
            ['products', 'supplier', 'relation', '→ suppliers', 'Поставщик (M:1)'],
            ['products', 'categories', 'relation multi', '→ categories', 'Категории (M:N)'],
            ['products', 'deleted', 'date', 'nullable', 'Логическое удаление'],
            ['customers', 'full_name', 'text', 'обязательное', 'ФИО'],
            ['customers', 'email', 'email', '', 'Почта'],
            ['loyalty_cards', 'customer', 'relation', 'уникальное', 'Связь 1:1'],
            ['loyalty_cards', 'tier', 'select', 'bronze/silver/gold', 'Уровень'],
            ['orders', 'status', 'select', 'pending…cancelled', 'Статус заказа'],
            ['orders', 'total', 'number', '', 'Итоговая сумма'],
            ['order_items', 'quantity', 'number', '≥ 1', 'Количество'],
            ['promotions', 'discount_percent', 'number', '0–100', 'Скидка %'],
            ['promotions', 'products', 'relation multi', '→ products', 'Товары акции'],
            ['reviews', 'rating', 'number', '1–5', 'Оценка'],
            ['users', 'role', 'select', 'customer/manager/admin', 'Роль'],
        ],
    )

    add_heading(doc, '5. Экспорт схемы', 1)
    add_para(
        doc,
        'Схема разворачивается скриптом pocketbase/setup.mjs (создание коллекций, API Rules, '
        'seed, демо-пользователи). Экспорт коллекций: pocketbase/pb_schema_export.json. '
        'Hook остатков: pocketbase/pb_hooks/orders.pb.js (HTTP 409 при нехватке stock).',
    )

    add_heading(doc, '6. Контракт API', 1)
    add_para(doc, 'Базовый адрес: {API_BASE_URL}, по умолчанию http://127.0.0.1:8090')
    add_table(
        doc,
        ['Метод', 'Путь', 'Назначение'],
        [
            ['POST', '/api/collections/users/auth-with-password', 'Вход'],
            ['POST', '/api/collections/users/auth-refresh', 'Обновление токена'],
            ['POST', '/api/collections/users/records', 'Регистрация'],
            ['GET', '/api/collections/{col}/records', 'Список (page, perPage, sort, filter, expand)'],
            ['GET', '/api/collections/{col}/records/{id}', 'Карточка'],
            ['POST', '/api/collections/{col}/records', 'Создание'],
            ['PATCH', '/api/collections/{col}/records/{id}', 'Изменение / soft-delete'],
            ['DELETE', '/api/collections/{col}/records/{id}', 'Физическое удаление'],
        ],
    )
    add_para(doc, 'Коды ошибок: 400/422 — валидация полей; 401 — нет/истёк токен; 403 — нет прав; '
             '404 — не найдено; 409 — конфликт (остаток); 5xx — сервер; network/timeout — недоступность.')
    add_para(doc, 'Ответ списка: { page, perPage, totalItems, totalPages, items: [...] }.')

    add_heading(doc, '7. Архитектура клиента', 1)
    add_para(
        doc,
        'Слои: models → repositories (интерфейс + PocketBase) → state (ChangeNotifier/provider) → '
        'screens/widgets. Экраны не обращаются к Dio напрямую. Адрес API задаётся --dart-define=API_BASE_URL.',
    )
    add_para(doc, 'Основные маршруты: /, /products, /categories, /brands, /suppliers, /customers, '
             '/orders, /promotions, /reviews, /checkout, /my-orders, /fulfillment, /analytics, /users, '
             '/login, /register, /forbidden.')
    add_table(
        doc,
        ['Раздел', 'customer', 'manager', 'admin'],
        [
            ['Каталог (просмотр)', 'да', 'да', 'да'],
            ['CRUD каталога', 'нет', 'да', 'да'],
            ['Мои заказы / checkout', 'да (уник.)', 'нет', 'нет'],
            ['Сборка заказов', 'нет', 'да (уник.)', 'да'],
            ['Пользователи / аналитика', 'нет', 'нет', 'да (уник.)'],
            ['Hard delete / restore', 'нет', 'нет', 'да'],
        ],
    )
    add_para(doc, 'Состояния списков: loading, success, empty, error (виджет AsyncBody).')

    add_heading(doc, '8. Ключевые фрагменты реализации', 1)
    add_para(
        doc,
        'OrderPricingEngine (lib/features/pricing/order_pricing_engine.dart) — расчёт subtotal, '
        'скидок и доставки. PbCollectionRepository — общий CRUD PocketBase. AuthNotifier + '
        'GoRouter.redirect — защита маршрутов. Permissions — скрытие элементов UI. ApiClient — '
        'Authorization, auth-refresh при 401, map 400/422→ValidationException, 409→ConflictException.',
    )

    add_heading(doc, '9. Демонстрация работы', 1)
    add_para(
        doc,
        'Приложите снимки экранов: вход под каждой ролью; каталог с фильтрами в URL; форма товара '
        'со связями; оформление заказа с разбивкой цены; экран «Сборка»; аналитика; отказ /forbidden; '
        'ошибка при остановленном сервере. Учётные записи: admin@fitshop.ru / manager@fitshop.ru / '
        'customer@fitshop.ru, пароль Fitshop123!',
    )
    add_para(doc, 'Репозиторий: (указать ссылку после публикации на GitHub).')
    add_para(doc, 'Опубликованное приложение: GitHub Pages (после push в main и настройки secret API_BASE_URL).')

    add_heading(doc, '10. Перечень тестов', 1)
    add_table(
        doc,
        ['Файл', 'Что проверяет'],
        [
            ['unit_test.dart', 'Validators: обязательность, email, пароль, положительное число'],
            ['unit_test.dart', 'Permissions: права customer/manager/admin и canAccessPath'],
            ['unit_test.dart', 'Product.fromJson: устойчивость к пустым полям и expand'],
            ['unit_test.dart', 'OrderPricingEngine: промо, gold, free shipping, пустые qty'],
            ['widget_states_test.dart', 'AsyncBody: loading / empty / error+retry / success'],
            ['widget_states_test.dart', 'Скрытие кнопки управления каталогом у customer'],
            ['widget_test.dart', 'Редирект неавторизованного пользователя на /login'],
        ],
    )
    add_para(doc, 'Запуск: flutter test — не менее 8 модульных и 5 виджет-тестов, все проходят.')

    add_heading(doc, '11. Вывод', 1)
    add_para(
        doc,
        'Разработан полноценный клиент магазина фитнес-продукции на Flutter Web с сервером PocketBase. '
        'Реализованы восемь+ связанных сущностей, CRUD, валидация, адаптивная вёрстка, аутентификация '
        'и три роли с уникальными разделами, серверные фильтры и пагинация, soft/hard delete, обработка '
        'сетевых ошибок и обновление токена. Содержательная особенность — расчёт стоимости заказа с '
        'акциями, лояльностью и доставкой. Проект покрыт тестами и подготовлен к публикации на GitHub Pages.',
    )

    doc.save(OUT)
    print('Saved', OUT)


if __name__ == '__main__':
    main()
